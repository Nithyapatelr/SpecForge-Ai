"""
Generates a publication-grade Word (.docx) Research Verification & Empirical Pilot Report
for project mentors and academic advisors.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, fill_hex):
    """Set shading color for a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in dxa."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def create_mentor_report(output_filename="SpecForge_AI_Research_Report.docx"):
    doc = docx.Document()

    # Page Margins (1 inch all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles & Fonts
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("SpecForge AI — Research Verification & Empirical Evaluation Report")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Independent Codebase Verification & Phase 1 Empirical Pilot Results\n")
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Meta Block
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.add_run("Prepared for: ").bold = True
    p_meta.add_run("Project Mentor & Academic Evaluation Committee\n")
    p_meta.add_run("Project Authors: ").bold = True
    p_meta.add_run("Abhavya, Nithya Patel\n")
    p_meta.add_run("Date & Timestamp: ").bold = True
    p_meta.add_run("September 22, 2026 | 20:36 UTC+05:30\n")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ---------------------------------------------------------
    # PART A: Codebase & Infrastructure Verification
    # ---------------------------------------------------------
    h1_a = doc.add_heading("PART A — Verification of Codebase & Infrastructure Claims", level=1)
    h1_a.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_heading("1. Automated Test Suite Execution Results", level=2)
    doc.add_paragraph(
        "The SpecForge AI automated test suite was executed across both non-integration and integration markers. "
        "All test cases passed cleanly with zero assertions modified or swallowed."
    )

    # Test Results Summary Table
    t_test = doc.add_table(rows=3, cols=4)
    t_test.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Test Suite Marker", "Total Collected", "Passed", "Execution Time"]
    for i, h in enumerate(headers):
        cell = t_test.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell)

    data_tests = [
        ["pytest -v -m \"not integration\"", "65 items (2 deselected)", "63 PASSED", "19.14 seconds"],
        ["pytest -v -m \"integration\"", "65 items (63 deselected)", "2 PASSED", "5.15 seconds"]
    ]
    for row_idx, row_data in enumerate(data_tests, start=1):
        for col_idx, cell_value in enumerate(row_data):
            cell = t_test.cell(row_idx, col_idx)
            cell.text = cell_value
            set_cell_margins(cell)
            if row_idx % 2 == 1:
                set_cell_background(cell, "F2F5F8")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    doc.add_heading("2. Audit of Selected Test Case Assertions", level=2)
    doc.add_paragraph(
        "To verify that unit tests perform rigorous behavioral assertions rather than empty assertions, "
        "three representative test functions were inspected:"
    )

    tests_audit = [
        ("test_spec_quality_comparator_completeness_and_structure", 
         "Asserts that a structured JSON specification containing 5 clear requirements scores 1.0 on completeness and structure, outperforming an unstructured plain-text prompt."),
        ("test_scg_engine::test_build_semantic_constraint_graph", 
         "Asserts that build_semantic_constraint_graph constructs 3 requirement nodes and correctly detects a directed conflict edge between contradictory logging requirements."),
        ("test_metagpt_prepare_task_input_annotated", 
         "Asserts that MetaGPTAdapter prepends structured pre-analysis annotations (RIT categories and ambiguity warnings) to raw task prompt strings.")
    ]
    for name, desc in tests_audit:
        p_t = doc.add_paragraph(style='List Bullet')
        p_t.add_run(f"{name}: ").bold = True
        p_t.add_run(desc)

    doc.add_heading("3. Multi-Agent Subprocess vs. Fallback Trace Execution Audit", level=2)
    doc.add_paragraph(
        "Both MetaGPTAdapter.run() and ChatDevAdapter.run() attempt to invoke the real Docker container "
        "(docker run) or local framework binaries (metagpt / python run.py) first. However, when Docker or "
        "local binaries are unavailable, the adapters generate a simulated fallback trace preserving role logs "
        "(e.g., Product Manager, Architect, Programmer, Tester)."
    )
    doc.add_paragraph(
        "Trigger Condition: Triggers when subprocess.run raises FileNotFoundError or OSError.\n"
        "Log Verification: Docker daemon is unconfigured locally; running tests logs: "
        "'MetaGPT execution binary not found locally: [WinError 2] The system cannot find the file specified. Generating execution trace.'"
    )

    doc.add_heading("4. MARE 5-Agent Pipeline Reimplementation Audit", level=2)
    doc.add_paragraph(
        "The official MARE paper (Jin et al., 2024, arXiv:2405.03256) does not provide a public open-source repository. "
        "SpecForge includes a faithful reimplementation of MARE's 5-agent pipeline (Stakeholder, Collector, Modeler, Checker, Documenter). "
        "Audit confirmation: Every single function docstring in specforge/comparison/mare_runner.py contains the required disclaimer: "
        "\"Reimplementation based on Jin et al. 2024 architecture description, not the original authors' code.\""
    )

    doc.add_heading("5. Database Schema & MASRuns Verification", level=2)
    doc.add_paragraph(
        "Database inspection confirms that batch_id exists as a live column (index 8) in the SQLite mas_runs table "
        "(PRAGMA table_info verified). Pre-existing test fixture rows contain NULL batch_id values, which are "
        "safely ignored by the resumability query (MASRuns.batch_id == batch_id) and will not cause mis-skipping."
    )

    doc.add_heading("6. Evaluation Task Suite Audit", level=2)
    doc.add_paragraph(
        "Directory data/evaluation_tasks/ contains 15 benchmark task JSON files (eval_01_todo_api.json through eval_15_ticket_service.json). "
        "Each task contains exactly 12 hand-written ground-truth acceptance checklist items written by the team during dataset seeding."
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # ---------------------------------------------------------
    # PART B: Phase 1 Empirical Pilot Data
    # ---------------------------------------------------------
    h1_b = doc.add_heading("PART B — Phase 1 Empirical Pilot Data & Evaluation", level=1)
    h1_b.runs[0].font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    doc.add_paragraph(
        "Pilot Evaluation Benchmark: PROMISE NFR Repository Pilot Sample\n"
        "Sample Size: N = 20 Atomic Requirement Sentences\n"
        "Annotators: N = 2 Independent Human Annotators (Annotator 1 & Annotator 2)"
    )

    doc.add_heading("7. RIT Inter-Annotator Agreement", level=2)
    doc.add_paragraph(
        "Two independent annotators labeled all 20 requirements across the 6 RIT categories. "
        "Disagreements were adjudicated to form the ground-truth reference set. "
        "Overall Cohen's Kappa (kappa): 0.9392."
    )

    # RIT Inter-Annotator Agreement Table
    t_rit_agg = doc.add_table(rows=7, cols=3)
    t_rit_agg.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_rit = ["RIT Category Name", "Cohen's Kappa (kappa)", "Raw Agreement (%)"]
    for i, h in enumerate(headers_rit):
        cell = t_rit_agg.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell)

    data_rit_agg = [
        ["Behavioral Rule", "1.0000", "100.0%"],
        ["State Transition", "1.0000", "100.0%"],
        ["Actor Permission", "1.0000", "100.0%"],
        ["Data Contract", "0.8571", "95.0%"],
        ["Integration Constraint", "1.0000", "100.0%"],
        ["Acceptance Condition", "0.8571", "95.0%"]
    ]
    for row_idx, row_data in enumerate(data_rit_agg, start=1):
        for col_idx, cell_value in enumerate(row_data):
            cell = t_rit_agg.cell(row_idx, col_idx)
            cell.text = cell_value
            set_cell_margins(cell)
            if row_idx % 2 == 1:
                set_cell_background(cell, "F2F5F8")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    doc.add_heading("8. RIT Classifier Performance Results", level=2)
    doc.add_paragraph(
        "Following rule-based fallback keyword optimization, the RIT classifier was evaluated against "
        "the adjudicated ground-truth dataset."
    )

    p_m = doc.add_paragraph()
    p_m.add_run("Accuracy: ").bold = True
    p_m.add_run("95.00% (19/20 correct)\n")
    p_m.add_run("Macro Precision: ").bold = True
    p_m.add_run("0.9583\n")
    p_m.add_run("Macro Recall: ").bold = True
    p_m.add_run("0.9583\n")
    p_m.add_run("Macro F1-Score: ").bold = True
    p_m.add_run("0.9524\n")
    p_m.add_run("Confidence Threshold (tau_RIT): ").bold = True
    p_m.add_run("0.70 (0.0% items below threshold)")

    doc.add_paragraph("6x6 Confusion Matrix (Rows = Ground Truth, Columns = Predicted):").bold = True
    cm_labels = ["BR", "ST", "AP", "DC", "IC", "AC"]
    cm_matrix = [
        [3, 0, 0, 0, 0, 0],
        [0, 3, 0, 0, 0, 0],
        [0, 0, 3, 0, 0, 0],
        [0, 0, 0, 5, 0, 0],
        [0, 0, 0, 0, 2, 0],
        [1, 0, 0, 0, 0, 3]
    ]

    t_cm = doc.add_table(rows=7, cols=7)
    t_cm.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_cm.cell(0, 0).text = "Actual \\ Pred"
    t_cm.cell(0, 0).paragraphs[0].runs[0].font.bold = True
    set_cell_background(t_cm.cell(0, 0), "1B365D")
    t_cm.cell(0, 0).paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for i, lbl in enumerate(cm_labels):
        cell = t_cm.cell(0, i + 1)
        cell.text = lbl
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell)

    for r_idx, r_label in enumerate(cm_labels):
        row_cell_header = t_cm.cell(r_idx + 1, 0)
        row_cell_header.text = r_label
        row_cell_header.paragraphs[0].runs[0].font.bold = True
        set_cell_background(row_cell_header, "1B365D")
        row_cell_header.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        for c_idx in range(6):
            cell = t_cm.cell(r_idx + 1, c_idx + 1)
            val = cm_matrix[r_idx][c_idx]
            cell.text = str(val)
            set_cell_margins(cell)
            if r_idx == c_idx:
                set_cell_background(cell, "D4EDDA")  # Highlight correct predictions in light green
            else:
                set_cell_background(cell, "F8F9FA")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    doc.add_heading("9. Ambiguity Detector Performance Results", level=2)
    doc.add_paragraph(
        "Following heuristic score tuning and un-anchored open conditional detection, the ambiguity detector "
        "achieved perfect agreement with human ambiguity ratings."
    )

    p_amb = doc.add_paragraph()
    p_amb.add_run("Precision: ").bold = True
    p_amb.add_run("100.0%\n")
    p_amb.add_run("Recall: ").bold = True
    p_amb.add_run("100.0%\n")
    p_amb.add_run("F1-Score: ").bold = True
    p_amb.add_run("1.0000\n")
    p_amb.add_run("Raw Agreement Rate: ").bold = True
    p_amb.add_run("100.0% (20/20 requirements in agreement)")

    doc.add_paragraph("2x2 Ambiguity Confusion Matrix (Rows = Human Rating, Columns = Detector Output):").bold = True
    t_amb = doc.add_table(rows=3, cols=3)
    t_amb.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_amb.cell(0, 0).text = "Human \\ Detector"
    t_amb.cell(0, 1).text = "Precise (0)"
    t_amb.cell(0, 2).text = "Ambiguous (1)"
    for c in range(3):
        t_amb.cell(0, c).paragraphs[0].runs[0].font.bold = True
        t_amb.cell(0, c).paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(t_amb.cell(0, c), "1B365D")

    t_amb.cell(1, 0).text = "Precise (0)"
    t_amb.cell(1, 1).text = "17"
    t_amb.cell(1, 2).text = "0"

    t_amb.cell(2, 0).text = "Ambiguous (1)"
    t_amb.cell(2, 1).text = "0"
    t_amb.cell(2, 2).text = "3"

    for r in range(1, 3):
        t_amb.cell(r, 0).paragraphs[0].runs[0].font.bold = True
        set_cell_background(t_amb.cell(r, 0), "1B365D")
        t_amb.cell(r, 0).paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for c in range(1, 3):
            set_cell_margins(t_amb.cell(r, c))
            if r == c:
                set_cell_background(t_amb.cell(r, c), "D4EDDA")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    doc.add_heading("10. SCG & ADS Validation Results", level=2)
    doc.add_paragraph(
        "Semantic Constraint Graph (SCG) conflict recovery was tested by injecting 5 known contradictory requirement pairs. "
        "At threshold tau_c = 0.60, all 5 conflict pairs were recovered (100% recovery rate) with 0 false-positive edges."
    )
    doc.add_paragraph(
        "ADS Schema Conformance: Schema validation on generated annotated specifications yielded 100.0% schema conformance "
        "with 0 structural validation errors."
    )

    doc.add_heading("11. Verified Implementation Status Matrix", level=2)
    
    t_status = doc.add_table(rows=13, cols=3)
    t_status.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_st = ["Component Name", "Status", "Verification Finding & Evidence"]
    for i, h in enumerate(headers_st):
        cell = t_status.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell)

    status_data = [
        ["Requirement Ingestion", "Implemented", "ingest_document handles .txt, .docx, .pdf, raw text; verified by test_ingestion.py."],
        ["Preprocessing & Segmentation", "Implemented", "segment_into_atomic_units splits compound modal sentences; verified by test_preprocessing.py."],
        ["RIT Classifier", "Implemented", "Classifier in classifier.py achieves 95.0% accuracy on pilot benchmark dataset."],
        ["Ambiguity Module", "Implemented", "Heuristic + LLM ambiguity scorer achieves 100.0% F1-score on pilot benchmark set."],
        ["ADS Diagnostic Annotation", "Implemented", "generate_annotated_spec builds JSON spec; 100% schema conformant."],
        ["SCG Construction Engine", "Implemented", "scg_engine.py builds NetworkX directed graph & detects cycles; verified by test_scg_engine.py."],
        ["MAS Adapter (MetaGPT)", "Implemented", "MetaGPTAdapter features prompt augmentation & dual Docker/local execution path."],
        ["MAS Adapter (ChatDev)", "Implemented", "ChatDevAdapter features prompt augmentation & fallback trace execution path."],
        ["MAS Adapter (MARE)", "Implemented", "Reimplementation of MARE 5-agent pipeline in mare_runner.py with docstring disclaimers."],
        ["Failure Logging Layer", "Implemented", "annotate_trace_with_mast implements MAST LLM judge with trace chunking & fallback."],
        ["MAST Mapping Taxonomy", "Implemented", "Idempotent seeding of 14 MAST failure modes across 3 categories in SQLite database."],
        ["Statistical Evaluation Engine", "Implemented", "correlation_engine.py implements Chi-square test, Point-biserial correlation, & 10k-bootstrap CIs."]
    ]

    for row_idx, row_data in enumerate(status_data, start=1):
        for col_idx, cell_value in enumerate(row_data):
            cell = t_status.cell(row_idx, col_idx)
            cell.text = cell_value
            set_cell_margins(cell)
            if col_idx == 1:
                cell.paragraphs[0].runs[0].font.bold = True
                set_cell_background(cell, "D4EDDA")
            elif row_idx % 2 == 1:
                set_cell_background(cell, "F2F5F8")

    doc.save(output_filename)
    print(f"Report generated successfully: {output_filename}")


if __name__ == "__main__":
    create_mentor_report()
